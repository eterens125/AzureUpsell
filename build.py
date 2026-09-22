#!/usr/bin/env python3
"""Build the Azure Upsell Leaderboard page from the Domino's "UPSELL BY TM" reports.

Usage:  python3 build.py            # reads ./data, writes ./dist/index.html

Inputs (drop new files in ./data):
  * a multi-sheet period workbook  - one sheet per coupon code (sheet name = code)
  * single-day reports in data/daily/ - one file per coupon, named  YYYY-MM-DD_<code>.xlsx
    (single sheet "repUpsell"; the coupon code is taken from the file name)
Every report carries its own Begin/End Date in the header; overlapping dates are refused
so a day can't be counted twice.
"""
import glob, json, os, re, sys, datetime as dt
import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.abspath(__file__))
PERIOD_NAME = "Period 10 to date"
MIN_COUPON_ORDER = ["917526", "8255", "1126", "5385", "4342", "8210", "8253", "8225", "5073", "8666", "DIP0226"]


def nice(n):
    """'Mills, Timothy' -> 'Timothy M.' (first name + last initial only)."""
    if "," in n:
        last, first = [x.strip() for x in n.split(",", 1)]
        return f"{first} {last[:1].upper()}."
    return n


def header_dates(ws):
    text = " ".join(str(c) for r in ws.iter_rows(min_row=1, max_row=2, values_only=True) for c in r if c)
    b = re.search(r"Begin Date:\s*(\d+/\d+/\d+)", text)
    e = re.search(r"End Date:\s*(\d+/\d+/\d+)", text)
    if not (b and e):
        raise ValueError("no Begin/End Date in header")
    p = lambda s: dt.datetime.strptime(s, "%m/%d/%Y").date()
    return p(b.group(1)), p(e.group(1))


def parse_sheet(ws, coupon):
    rows, store = [], None
    for r in ws.iter_rows(min_row=5, values_only=True):
        name, code, prod, qty, ordr, tm = r[1], r[3], r[5], r[6], r[7], r[8]
        if name and prod in ("", None) and qty in ("", None):
            store = name                      # store header row
            continue
        if name in ("", None):
            continue                          # store total row (checked below)
        if name and prod:
            rows.append(dict(coupon=coupon, store=store, name=name, code=code or "",
                             ord=float(ordr or 0), tm=float(tm or 0), internet=(name == "Internet")))
    return rows


def load_sources():
    rows, sources = [], []
    files = sorted(glob.glob(os.path.join(ROOT, "data", "*.xlsx")) + glob.glob(os.path.join(ROOT, "data", "daily", "*.xlsx")))
    if not files:
        sys.exit("No .xlsx files found in ./data")
    for f in files:
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            coupon = (re.split(r"[_-]", os.path.splitext(os.path.basename(f))[0])[-1] if ws.title == "repUpsell" else ws.title).upper()
            start, end = header_dates(ws)
            part = parse_sheet(ws, coupon)
            # sanity: rows must equal the report's own store total rows
            tot = {}
            store = None
            for r in ws.iter_rows(min_row=5, values_only=True):
                if r[1] and r[5] in ("", None) and r[6] in ("", None):
                    store = r[1]
                elif r[1] in ("", None) and r[6] not in ("", None):
                    tot[store] = (float(r[7] or 0), float(r[8] or 0))
            for s, (o, t) in tot.items():
                po = sum(x["ord"] for x in part if x["store"] == s)
                pt = sum(x["tm"] for x in part if x["store"] == s)
                if (po, pt) != (o, t):
                    raise ValueError(f"{os.path.basename(f)} [{coupon}] store {s}: rows ({po},{pt}) != report total ({o},{t})")
            for x in part:
                x["start"], x["end"] = start, end
            rows += part
            sources.append((start, end, os.path.basename(f), coupon))
    return pd.DataFrame(rows), sources


def merge_ranges(ranges):
    out = []
    for s, e in sorted(set(ranges)):
        if out and s <= out[-1][1] + dt.timedelta(days=1):
            out[-1][1] = max(out[-1][1], e)
        else:
            out.append([s, e])
    return out


def check_overlap(sources):
    per = {}
    for s, e, f, c in sources:
        per.setdefault(c, []).append((s, e, f))
    for c, lst in per.items():
        lst.sort()
        for a, b in zip(lst, lst[1:]):
            if b[0] <= a[1]:
                sys.exit(f"Coupon {c}: {a[2]} ({a[0]}..{a[1]}) overlaps {b[2]} ({b[0]}..{b[1]}); a day would be counted twice.")


def fmt_range(s, e):
    d = lambda x: x.strftime("%b ") + str(x.day)
    if s == e:
        return d(s)
    return f"{d(s)} &ndash; {d(e)}" if s.month != e.month else f"{d(s)} &ndash; {e.day}"


def build_data(df):
    coupons = MIN_COUPON_ORDER + sorted(set(df.coupon) - set(MIN_COUPON_ORDER))
    stores = sorted(df.store.unique())
    tm, net = df[~df.internet], df[df.internet]
    ref = "5073" if "5073" in set(df.coupon) else sorted(set(df.coupon))[0]  # every sheet repeats each person's TM Ord
    S = []
    for s in stores:
        t, n = tm[tm.store == s], net[net.store == s]
        by = {}
        for c in coupons:
            g = t[(t.coupon == c) & (t.ord > 0)].groupby("name").ord.sum().reset_index().sort_values(["ord", "name"], ascending=[False, True])
            by[c] = dict(tm=int(t[t.coupon == c].ord.sum()), net=int(n[n.coupon == c].ord.sum()),
                         top=[dict(n=nice(r["name"]), v=int(r["ord"])) for _, r in g.iterrows()])
        team = []
        for (name, code), g in t.groupby(["name", "code"]):
            tmo = int(g[g.coupon == ref].tm.sum())
            if tmo > 0:
                team.append(dict(n=nice(name), o=int(g.ord.sum()), t=tmo))
        st = dict(id=s, tmOrd=int(t[t.coupon == ref].tm.sum()), ups=int(t.ord.sum()), netOrd=int(n.ord.sum()),
                  netTm=int(n[n.coupon == ref].tm.sum()), by=by, team=team)
        assert sum(x["o"] for x in team) == st["ups"] and sum(x["t"] for x in team) == st["tmOrd"], s
        S.append(st)
    people = []
    for (name, code), g in tm.groupby(["name", "code"]):
        tmo = int(g[g.coupon == ref].tm.sum())
        if tmo > 0:
            people.append(dict(name=nice(name), stores=sorted(g.store.unique()), ord=int(g.ord.sum()), tm=tmo))
    return dict(stores=S, coupons=coupons, people=people)


def main():
    df, sources = load_sources()
    check_overlap(sources)
    data = build_data(df)
    ranges = merge_ranges([(s, e) for s, e, _, _ in sources])
    period = " plus ".join(fmt_range(s, e) for s, e in ranges) + f", {ranges[-1][1].year}"
    span = {}
    for s, e, f, c in sources:
        span.setdefault(f, (s, e))
    multi = sorted({(s, e) for f, (s, e) in span.items() if s != e})
    single = sorted({s for f, (s, e) in span.items() if s == e})
    parts = []
    if multi:
        parts.append("the UPSELL BY TM report (" + ", ".join(fmt_range(s, e) for s, e in multi) + ")")
    if single:
        parts.append("single-day coupon reports for " + ", ".join(fmt_range(d, d) for d in single))
    sources_txt = " plus ".join(parts)
    html = open(os.path.join(ROOT, "template.html"), encoding="utf-8").read()
    for k, v in {"{{PERIOD_NAME}}": PERIOD_NAME, "{{PERIOD}}": period, "{{SOURCES}}": sources_txt}.items():
        assert k in html, k
        html = html.replace(k, v)
    assert "/*DATA*/" in html
    html = html.replace("/*DATA*/", json.dumps(data, separators=(",", ":")))
    # template.html is a bare content fragment (no doctype/html/head/body) by design,
    # since the Claude Artifact publish flow supplies its own page skeleton with a
    # UTF-8 charset at publish time. This standalone GitHub build has no such wrapper,
    # so without an explicit <meta charset> a browser opening dist/index.html directly
    # (or via a host that doesn't send a charset header) can guess the wrong encoding
    # and mangle multi-byte characters like the em dash into "â€”" mojibake. Wrap it in
    # a real document here so the charset is always declared.
    html = (
        "<!DOCTYPE html>\n"
        '<html lang="en">\n<head>\n<meta charset="UTF-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"{html}\n</head>\n<body></body>\n</html>\n"
    )
    out = os.path.join(ROOT, "dist", "index.html")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    open(out, "w", encoding="utf-8").write(html)
    ups = sum(s["ups"] for s in data["stores"]); tmo = sum(s["tmOrd"] for s in data["stores"])
    print(f"Built {out}\n  period: {period}\n  team-member upsell orders: {ups} of {tmo} ({ups / tmo:.1%})")


if __name__ == "__main__":
    main()
